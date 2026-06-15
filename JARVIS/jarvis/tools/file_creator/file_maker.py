"""
jarvis/tools/file_creator/file_maker.py

Creates PDF, PPTX, DOCX, XLSX files on command.
All files are saved to C:/Users/{user}/Documents/JARVIS_Files/ by default.
"""

import os
from pathlib import Path
from datetime import datetime


# ── Output directory ───────────────────────────────────────────
def get_output_dir() -> Path:
    username = os.getenv("USERNAME", "User")
    out_dir = Path(f"C:/Users/{username}/Documents/JARVIS_Files")
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir

def _timestamped(name: str, ext: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = "".join(c if c.isalnum() or c in " _-" else "_" for c in name)
    return get_output_dir() / f"{safe_name}_{ts}.{ext}"


# ── PDF ────────────────────────────────────────────────────────
def create_pdf(title: str, content: str, filename: str = None) -> str:
    """Create a styled PDF with title and content."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.enums import TA_CENTER

        out_path = Path(filename) if filename else _timestamped(title, "pdf")
        doc = SimpleDocTemplate(str(out_path), pagesize=A4,
                                leftMargin=2.5*cm, rightMargin=2.5*cm,
                                topMargin=2.5*cm, bottomMargin=2.5*cm)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "JARVISTitle",
            fontSize=20, fontName="Helvetica-Bold",
            spaceAfter=12, alignment=TA_CENTER,
            textColor=colors.HexColor("#1a1a2e")
        )
        body_style = ParagraphStyle(
            "JARVISBody",
            fontSize=11, fontName="Helvetica",
            spaceAfter=8, leading=16
        )

        story = [
            Paragraph(title, title_style),
            HRFlowable(width="100%", thickness=2, color=colors.HexColor("#e94560")),
            Spacer(1, 0.5*cm),
        ]

        for para in content.split("\n\n"):
            para = para.strip()
            if not para:
                continue
            if para.startswith("# "):
                story.append(Paragraph(para[2:], styles["Heading1"]))
            elif para.startswith("## "):
                story.append(Paragraph(para[3:], styles["Heading2"]))
            else:
                story.append(Paragraph(para.replace("\n", "<br/>"), body_style))
            story.append(Spacer(1, 0.2*cm))

        doc.build(story)
        return str(out_path)
    except Exception as e:
        return f"PDF creation failed: {e}"


# ── PPTX ───────────────────────────────────────────────────────
def create_pptx(title: str, slides: list, filename: str = None) -> str:
    """
    Create a PowerPoint presentation.
    slides = [{"title": "...", "content": "bullet1\nbullet2"}, ...]
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN

        prs = Presentation()
        prs.slide_width  = Inches(13.33)
        prs.slide_height = Inches(7.5)

        # Title slide
        title_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = title
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
        slide.shapes.title.text_frame.paragraphs[0].font.bold = True
        if slide.placeholders[1]:
            slide.placeholders[1].text = f"Created by JARVIS — {datetime.now().strftime('%d %b %Y')}"

        # Content slides
        content_layout = prs.slide_layouts[1]
        for s in slides:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = s.get("title", "")
            slide.shapes.title.text_frame.paragraphs[0].font.bold = True
            slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)

            tf = slide.placeholders[1].text_frame
            tf.clear()
            lines = s.get("content", "").split("\n")
            for i, line in enumerate(lines):
                if not line.strip():
                    continue
                para = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                para.text = line.lstrip("•-* ")
                para.font.size = Pt(18)
                para.level = 0

        out_path = Path(filename) if filename else _timestamped(title, "pptx")
        prs.save(str(out_path))
        return str(out_path)
    except Exception as e:
        return f"PPTX creation failed: {e}"


# ── DOCX ───────────────────────────────────────────────────────
def create_docx(title: str, content: str, filename: str = None) -> str:
    """Create a Word document."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title_para.runs[0]
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

        doc.add_paragraph("")  # spacer

        # Body content
        for block in content.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("# "):
                doc.add_heading(block[2:], level=1)
            elif block.startswith("## "):
                doc.add_heading(block[3:], level=2)
            elif block.startswith("### "):
                doc.add_heading(block[4:], level=3)
            else:
                para = doc.add_paragraph(block)
                para.style.font.size = Pt(11)

        # Footer
        section = doc.sections[0]
        footer = section.footer
        footer.paragraphs[0].text = f"Generated by JARVIS — {datetime.now().strftime('%d %b %Y')}"
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        out_path = Path(filename) if filename else _timestamped(title, "docx")
        doc.save(str(out_path))
        return str(out_path)
    except Exception as e:
        return f"DOCX creation failed: {e}"


# ── XLSX ───────────────────────────────────────────────────────
def create_xlsx(title: str, data: list, headers: list = None, filename: str = None) -> str:
    """
    Create an Excel spreadsheet.
    data = [[row1col1, row1col2, ...], [row2col1, ...]]
    headers = ["Col1", "Col2", ...] (optional)
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = title[:31]   # Excel sheet name limit

        # Header style
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill("solid", fgColor="1a1a2e")
        header_align = Alignment(horizontal="center", vertical="center")

        row_start = 1
        if headers:
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font   = header_font
                cell.fill   = header_fill
                cell.alignment = header_align
            row_start = 2

        # Data rows
        alt_fill = PatternFill("solid", fgColor="f0f0f0")
        for row_idx, row in enumerate(data, row_start):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

        # Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

        out_path = Path(filename) if filename else _timestamped(title, "xlsx")
        wb.save(str(out_path))
        return str(out_path)
    except Exception as e:
        return f"XLSX creation failed: {e}"


# ── Unified tool function ──────────────────────────────────────
def create_file_tool(file_type: str, title: str, content_or_data, **kwargs) -> str:
    """
    Single entry point for all file creation.
    file_type: "pdf" | "pptx" | "docx" | "xlsx"
    """
    file_type = file_type.lower().strip(".")
    if file_type == "pdf":
        return create_pdf(title, content_or_data, kwargs.get("filename"))
    elif file_type in ("pptx", "ppt", "presentation", "slides"):
        slides = content_or_data if isinstance(content_or_data, list) else [
            {"title": "Overview", "content": content_or_data}
        ]
        return create_pptx(title, slides, kwargs.get("filename"))
    elif file_type in ("docx", "doc", "word"):
        return create_docx(title, content_or_data, kwargs.get("filename"))
    elif file_type in ("xlsx", "xls", "excel", "spreadsheet"):
        headers = kwargs.get("headers", [])
        data = content_or_data if isinstance(content_or_data, list) else []
        return create_xlsx(title, data, headers, kwargs.get("filename"))
    else:
        return f"Unknown file type: {file_type}. Supported: pdf, pptx, docx, xlsx"
